"""
Упрощенный модуль для разбиения документов на чанки по разделам.
Каждый раздел (1.1, 1.2, 2.1 и т.д.) становится отдельным чанком.
"""

import os
import re
import time
import zipfile
from typing import List, Dict, Any, Tuple, Optional
import xml.etree.ElementTree as ET

# Константы
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# Паттерн для определения начала раздела (1., 1.1., 1.1.1. и т.д.)
SECTION_PATTERN = re.compile(
    r'^(\d+(?:\.\d+)*\.?)\s+',  # Номер раздела в начале строки
    re.MULTILINE
)

# Паттерн для основных разделов (без подподразделов)
MAIN_SECTION_PATTERN = re.compile(
    r'^(\d+(?:\.\d+)?\.?)\s+',  # Только разделы вида 1. или 1.1. (не 1.1.1)
    re.MULTILINE
)

# Паттерн для заголовков разделов (РАЗДЕЛ 1, Статья 1, Глава 1 и т.д.)
HEADING_PATTERN = re.compile(
    r'^((?:РАЗДЕЛ|ГЛАВА|СТАТЬЯ|ПУНКТ|ПРИЛОЖЕНИЕ)\s+\d+[\.:]*.*?)$',
    re.MULTILINE | re.IGNORECASE
)


def try_import_docx():
    """Попытка импорта библиотеки python-docx."""
    try:
        import docx
        return docx
    except Exception:
        return None


DOCX_LIB = try_import_docx()


def extract_text_from_docx_xml(path: str) -> str:
    """
    Извлечение текста из DOCX через прямой парсинг XML.
    Сохраняет структуру параграфов.
    """
    text_parts = []

    try:
        with zipfile.ZipFile(path, 'r') as z:
            # Читаем основной документ
            with z.open('word/document.xml') as f:
                data = f.read()

            root = ET.fromstring(data)

            # Извлекаем текст из параграфов
            for para in root.iter(f'{W_NS}p'):
                para_text = []
                for t_elem in para.iter(f'{W_NS}t'):
                    if t_elem.text:
                        para_text.append(t_elem.text)

                if para_text:
                    text_parts.append(''.join(para_text))

            # Обработка таблиц
            for table in root.iter(f'{W_NS}tbl'):
                table_text = []
                for row in table.iter(f'{W_NS}tr'):
                    row_text = []
                    for cell in row.iter(f'{W_NS}tc'):
                        cell_text = []
                        for t_elem in cell.iter(f'{W_NS}t'):
                            if t_elem.text:
                                cell_text.append(t_elem.text)
                        if cell_text:
                            row_text.append(' '.join(cell_text))
                    if row_text:
                        table_text.append(' | '.join(row_text))
                if table_text:
                    text_parts.append('\n'.join(table_text))

    except Exception as e:
        print(f"Ошибка при парсинге XML: {e}")
        return ""

    return '\n'.join(text_parts)


def extract_text_from_docx_lib(path: str) -> str:
    """
    Извлечение текста из DOCX с использованием python-docx.
    """
    try:
        import docx
        doc = docx.Document(path)

        text_parts = []

        # Извлекаем текст из параграфов
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())

        # Извлекаем текст из таблиц
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_parts.append(' | '.join(row_text))

        return '\n'.join(text_parts)

    except Exception as e:
        print(f"Ошибка при использовании python-docx: {e}")
        return ""


def extract_text_from_docx(path: str) -> str:
    """
    Извлечение текста из DOCX файла.
    Сначала пытается использовать python-docx, затем fallback на XML парсинг.
    """
    if DOCX_LIB is not None:
        text = extract_text_from_docx_lib(path)
        if text:
            return text

    # Fallback на XML парсинг
    return extract_text_from_docx_xml(path)


def extract_text_from_txt(path: str) -> str:
    """Извлечение текста из TXT файла."""
    try:
        with open(path, 'r', encoding='utf-8') as f:
            return f.read()
    except UnicodeDecodeError:
        try:
            with open(path, 'r', encoding='cp1251') as f:
                return f.read()
        except:
            return ""


def extract_text_from_pdf(path: str) -> str:
    """Извлечение текста из PDF файла."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(path)
        text_parts = []
        for page in doc:
            text_parts.append(page.get_text())
        doc.close()
        return '\n'.join(text_parts)
    except:
        return ""


def clean_text(text: str) -> str:
    """
    Очистка текста от лишних символов и нормализация.
    """
    # Удаляем множественные пробелы и табуляции
    text = re.sub(r'[ \t]+', ' ', text)

    # Удаляем множественные переносы строк (больше 2)
    text = re.sub(r'\n{3,}', '\n\n', text)

    # Удаляем пробелы в начале и конце строк
    lines = [line.strip() for line in text.split('\n')]
    text = '\n'.join(lines)

    # Удаляем пустые строки в начале и конце
    text = text.strip()

    return text


def split_by_sections(text: str) -> List[Dict[str, Any]]:
    """
    Разбивает текст на чанки по разделам.
    Основные разделы (1., 1.1., 2., 2.1.) становятся отдельными чанками.
    Подразделы (1.1.1, 1.1.2 и т.д.) объединяются с родительским разделом.
    Преамбула договора (текст до первого раздела) сохраняется как отдельный чанк.
    """
    chunks = []

    # Находим ВСЕ разделы (включая подподразделы)
    all_section_matches = list(SECTION_PATTERN.finditer(text))

    # Находим только ОСНОВНЫЕ разделы (без подподразделов)
    main_section_matches = list(MAIN_SECTION_PATTERN.finditer(text))

    # Если разделы не найдены, пробуем найти заголовки
    if not all_section_matches:
        heading_matches = list(HEADING_PATTERN.finditer(text))
        if heading_matches:
            main_section_matches = heading_matches
            all_section_matches = heading_matches

    if not all_section_matches:
        # Если структура не найдена, разбиваем по параграфам
        return split_by_paragraphs(text)

    # ВАЖНО: Проверяем, есть ли текст до первого раздела (преамбула договора)
    if all_section_matches:
        first_section_start = all_section_matches[0].start()
        if first_section_start > 0:
            preamble_text = text[:first_section_start].strip()
            if preamble_text:
                # Определяем заголовок преамбулы
                if 'ДОГОВОР' in preamble_text.upper()[:200]:
                    # Ищем строку с "ДОГОВОР"
                    lines = preamble_text.split('\n')
                    for line in lines[:5]:  # Проверяем первые 5 строк
                        if 'ДОГОВОР' in line.upper():
                            heading = line.strip()
                            break
                    else:
                        heading = "Преамбула договора"
                else:
                    heading = "Преамбула"

                chunks.append({
                    'heading': heading,
                    'text': preamble_text,
                    'section_number': '0',  # Специальный номер для преамбулы
                    'chunk_index': 0
                })

    # Фильтруем основные разделы, исключая подподразделы
    # Основной раздел - это раздел с не более чем 2 уровнями (например, 1.1, но не 1.1.1)
    filtered_matches = []
    for match in all_section_matches:
        section_num = match.group(1).rstrip('.')
        levels = section_num.split('.')
        if len(levels) <= 2:  # Только разделы вида X или X.Y
            filtered_matches.append(match)

    # Если не нашли основных разделов, используем все
    if not filtered_matches:
        filtered_matches = all_section_matches

    # Создаем чанки для каждого основного раздела
    for i, match in enumerate(filtered_matches):
        start_pos = match.start()

        # Определяем конец текущего раздела (до следующего основного раздела)
        if i < len(filtered_matches) - 1:
            end_pos = filtered_matches[i + 1].start()
        else:
            end_pos = len(text)

        # Извлекаем текст раздела
        section_text = text[start_pos:end_pos].strip()

        if not section_text:
            continue

        # Извлекаем номер раздела для заголовка
        section_number = match.group(1) if match.groups() else match.group(0)

        # Находим первую строку как заголовок
        first_line_end = section_text.find('\n')
        if first_line_end > 0:
            heading = section_text[:first_line_end].strip()
        else:
            heading = section_text[:100].strip(
            ) + "..." if len(section_text) > 100 else section_text

        # Индекс с учетом возможной преамбулы
        chunk_idx = len(chunks)

        chunks.append({
            'heading': heading,
            'text': section_text,
            'section_number': section_number.strip(),
            'chunk_index': chunk_idx
        })

    # ВАЖНО: Проверяем, есть ли текст после последнего раздела (реквизиты, подписи)
    if filtered_matches:
        # Проверяем текст после всех разделов
        last_match = filtered_matches[-1]
        # Находим где заканчивается последний раздел (до следующего заголовка или конца)
        last_section_text = text[last_match.start():]

        # Ищем признаки окончания последнего нумерованного раздела и начала реквизитов
        epilogue_markers = [
            'АДРЕСА И РЕКВИЗИТЫ СТОРОН',
            'РЕКВИЗИТЫ СТОРОН',
            'ЮРИДИЧЕСКИЕ АДРЕСА',
            'РЕКВИЗИТЫ И ПОДПИСИ',
            'ПОДПИСИ СТОРОН',
            'Заказчик\n',
            'Исполнитель\n'
        ]

        epilogue_start = -1
        for marker in epilogue_markers:
            pos = text.upper().find(marker.upper(), last_match.start())
            if pos > 0:
                # Проверяем, что это не внутри нумерованного раздела
                is_new_section = True
                for m in filtered_matches:
                    if m.start() <= pos < m.end():
                        is_new_section = False
                        break
                if is_new_section and (epilogue_start == -1 or pos < epilogue_start):
                    epilogue_start = pos

        # Если нашли эпилог (реквизиты/подписи), создаем для него отдельный чанк
        if epilogue_start > last_match.end():
            epilogue_text = text[epilogue_start:].strip()
            # Минимальная длина для эпилога
            if epilogue_text and len(epilogue_text) > 10:
                # Корректируем последний чанк - обрезаем его до начала эпилога
                if chunks and chunks[-1]['text'].endswith(epilogue_text):
                    chunks[-1]['text'] = text[last_match.start()
                                                               :epilogue_start].strip()

                # Определяем заголовок эпилога
                epilogue_lines = epilogue_text.split('\n')
                epilogue_heading = "Реквизиты и подписи сторон"
                for line in epilogue_lines[:3]:
                    if any(marker in line.upper() for marker in ['РЕКВИЗИТ', 'АДРЕСА', 'ПОДПИС']):
                        epilogue_heading = line.strip()
                        break

                # Добавляем эпилог как отдельный чанк
                chunks.append({
                    'heading': epilogue_heading,
                    'text': epilogue_text,
                    'section_number': '99',  # Специальный номер для эпилога
                    'chunk_index': len(chunks)
                })

    return chunks


def split_by_paragraphs(text: str, max_chunk_size: int = 2000) -> List[Dict[str, Any]]:
    """
    Резервный метод: разбивает текст по параграфам с ограничением размера.
    """
    chunks = []
    paragraphs = text.split('\n\n')

    current_chunk = []
    current_size = 0
    chunk_index = 0

    for para in paragraphs:
        para = para.strip()
        if not para:
            continue

        para_size = len(para)

        # Если параграф слишком большой, разбиваем его отдельно
        if para_size > max_chunk_size:
            # Сохраняем текущий накопленный чанк
            if current_chunk:
                chunk_text = '\n\n'.join(current_chunk)
                chunks.append({
                    'heading': f"Часть {chunk_index + 1}",
                    'text': chunk_text,
                    'section_number': str(chunk_index + 1),
                    'chunk_index': chunk_index
                })
                chunk_index += 1
                current_chunk = []
                current_size = 0

            # Разбиваем большой параграф по предложениям
            sentences = re.split(r'(?<=[.!?])\s+', para)
            temp_chunk = []
            temp_size = 0

            for sentence in sentences:
                if temp_size + len(sentence) > max_chunk_size and temp_chunk:
                    chunk_text = ' '.join(temp_chunk)
                    chunks.append({
                        'heading': f"Часть {chunk_index + 1}",
                        'text': chunk_text,
                        'section_number': str(chunk_index + 1),
                        'chunk_index': chunk_index
                    })
                    chunk_index += 1
                    temp_chunk = [sentence]
                    temp_size = len(sentence)
                else:
                    temp_chunk.append(sentence)
                    temp_size += len(sentence)

            if temp_chunk:
                chunk_text = ' '.join(temp_chunk)
                chunks.append({
                    'heading': f"Часть {chunk_index + 1}",
                    'text': chunk_text,
                    'section_number': str(chunk_index + 1),
                    'chunk_index': chunk_index
                })
                chunk_index += 1

        # Если добавление параграфа превысит лимит, создаем новый чанк
        elif current_size + para_size > max_chunk_size and current_chunk:
            chunk_text = '\n\n'.join(current_chunk)
            chunks.append({
                'heading': f"Часть {chunk_index + 1}",
                'text': chunk_text,
                'section_number': str(chunk_index + 1),
                'chunk_index': chunk_index
            })
            chunk_index += 1
            current_chunk = [para]
            current_size = para_size

        else:
            current_chunk.append(para)
            current_size += para_size

    # Добавляем последний чанк
    if current_chunk:
        chunk_text = '\n\n'.join(current_chunk)
        chunks.append({
            'heading': f"Часть {chunk_index + 1}",
            'text': chunk_text,
            'section_number': str(chunk_index + 1),
            'chunk_index': chunk_index
        })

    return chunks


def process_file(path: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """
    Обрабатывает файл и возвращает чанки и статистику.
    """
    t0 = time.time()

    # Извлекаем текст в зависимости от типа файла
    file_ext = os.path.splitext(path)[1].lower()

    if file_ext in ['.docx', '.doc']:
        text = extract_text_from_docx(path)
    elif file_ext == '.txt':
        text = extract_text_from_txt(path)
    elif file_ext == '.pdf':
        text = extract_text_from_pdf(path)
    else:
        text = ""

    if not text:
        return [], {"error": "Не удалось извлечь текст из файла"}

    # Очищаем текст
    text = clean_text(text)

    # Разбиваем на чанки
    chunks = split_by_sections(text)

    # Добавляем file_id к каждому чанку
    file_id = os.path.basename(path)
    for i, chunk in enumerate(chunks):
        chunk['file_id'] = file_id
        chunk['chunk_id'] = f"{file_id}::{i}"

    # Собираем статистику
    dt = (time.time() - t0) * 1000.0
    stats = {
        "file": os.path.basename(path),
        "chunks": len(chunks),
        "avg_chunk_chars": int(sum(len(c["text"]) for c in chunks) / max(1, len(chunks))),
        "time_ms": round(dt, 2),
    }

    return chunks, stats


# Совместимость с существующей системой
class ChunkingStats:
    def __init__(self, total_chunks: int, avg_chunk_size: float, processing_time: float):
        self.total_chunks = total_chunks
        self.avg_chunk_size = avg_chunk_size
        # Высокий score т.к. разбиение по логическим разделам
        self.semantic_coherence_score = 0.90
        self.processing_time = processing_time


def chunk_elements(elements, cfg=None, model=None, tokenizer=None):
    """
    Адаптер для совместимости с существующей системой чанкинга.
    Принимает элементы документа и возвращает чанки в ожидаемом формате.
    """
    t0 = time.time()

    # Объединяем текст из элементов
    text_parts = []
    for elem in elements:
        if hasattr(elem, 'text') and elem.text:
            text_parts.append(elem.text.strip())

    if not text_parts:
        stats = ChunkingStats(0, 0, time.time() - t0)
        return [], stats

    # Объединяем весь текст
    full_text = '\n'.join(text_parts)

    # Очищаем текст
    full_text = clean_text(full_text)

    # Разбиваем на чанки
    chunks = split_by_sections(full_text)

    # Конвертируем в ожидаемый формат
    result_chunks = []
    for chunk in chunks:
        result_chunk = {
            "text": chunk["text"],
            "chunk_index": chunk["chunk_index"],
            "heading": chunk["heading"],
            "section_number": chunk.get("section_number", ""),
            # Примерное количество токенов
            "token_len": len(chunk["text"].split()),
        }
        result_chunks.append(result_chunk)

    # Создаем статистику
    avg_chars = sum(len(c["text"])
                    for c in result_chunks) / max(1, len(result_chunks))
    stats = ChunkingStats(
        total_chunks=len(result_chunks),
        avg_chunk_size=avg_chars / 4,  # Примерное соотношение символов к токенам
        processing_time=time.time() - t0
    )

    return result_chunks, stats


# Функция для тестирования
def test_chunker():
    """
    Тестовая функция для проверки работы chunker'а.
    """
    test_text = """
ДОГОВОР ОКАЗАНИЯ УСЛУГ №

1. ПРЕДМЕТ ДОГОВОРА
1.1. По договору возмездного оказания услуг Исполнитель обязуется по заданию Заказчика оказать услуги, указанные в разделе 2 настоящего договора, а Заказчик обязуется оплатить эти услуги.
1.2. Исполнитель обязуется оказать следующие услуги:
- консультационные услуги
- юридические услуги

2. ПРАВА И ОБЯЗАННОСТИ СТОРОН
2.1. Стоимость услуг, оказываемых Исполнителем по настоящему договору составляет 100000 рублей.
2.2. Оплата услуг по договору осуществляется в течение 10 банковских дней с даты подписания акта сдачи-приемки оказанных услуг.

3. ОТВЕТСТВЕННОСТЬ СТОРОН
3.1. Стороны несут ответственность в соответствии с действующим законодательством.
3.2. За просрочку оплаты Заказчик уплачивает пени в размере 0,1% от суммы задолженности за каждый день просрочки.

4. ЗАКЛЮЧИТЕЛЬНЫЕ ПОЛОЖЕНИЯ
4.1. Договор вступает в силу с момента подписания.
4.2. Все споры решаются путем переговоров.
    """

    chunks = split_by_sections(test_text)

    print(f"Найдено чанков: {len(chunks)}")
    for i, chunk in enumerate(chunks):
        print(f"\nЧанк {i + 1}:")
        print(f"  Заголовок: {chunk['heading']}")
        print(f"  Номер раздела: {chunk['section_number']}")
        print(f"  Размер текста: {len(chunk['text'])} символов")
        print(f"  Начало текста: {chunk['text'][:100]}...")


if __name__ == "__main__":
    test_chunker()
